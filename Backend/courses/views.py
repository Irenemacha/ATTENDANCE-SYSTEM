from datetime import datetime, time as clock_time
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db import transaction
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response

from accounts.permissions import IsAdminOrStaff, IsLecturer, is_staff_or_superuser
from courses.models import Course, LecturerCourse, StudentCourse, Timetable, Classroom, Subject

User = get_user_model()

TIMETABLE_COLUMNS = ["course_code", "lecturer", "day", "start_time", "end_time", "room"]

DAY_ORDER = {code: index for index, (code, _) in enumerate(Timetable.DAYS)}


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_day_code(value):
    text = cell_text(value)
    if not text:
        return None
    for code, label in Timetable.DAYS:
        if text.upper() == code or text.lower() == label.lower():
            return code
    return None


def parse_clock_time(value):
    if isinstance(value, clock_time):
        return value
    if isinstance(value, datetime):
        return value.time()
    text = cell_text(value)
    for time_format in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, time_format).time()
        except ValueError:
            continue
    return None


def validate_timetable_data(data):
    errors = []
    cleaned = {}

    course_code = cell_text(data.get("course_code"))
    lecturer_name = cell_text(data.get("lecturer"))
    room = cell_text(data.get("room"))

    course = None
    if not course_code:
        errors.append("course_code is required.")
    else:
        course = Course.objects.filter(code__iexact=course_code).first()
        if not course:
            errors.append(f"course_code '{course_code}' does not match an existing course.")

    lecturer = None
    if not lecturer_name:
        errors.append("lecturer is required.")
    else:
        lecturer = User.objects.filter(username__iexact=lecturer_name).first()
        if not lecturer:
            errors.append(f"lecturer '{lecturer_name}' does not match an existing user.")
        elif not lecturer.is_staff and not lecturer.groups.filter(name__iexact="Lecturer").exists():
            errors.append(f"user '{lecturer_name}' is not a lecturer.")

    day = parse_day_code(data.get("day"))
    if not day:
        errors.append("day must be one of MON, TUE, WED, THU, FRI, SAT, SUN or the full day name.")

    start_time = parse_clock_time(data.get("start_time"))
    if start_time is None:
        errors.append("start_time must be a time like HH:MM.")

    end_time = parse_clock_time(data.get("end_time"))
    if end_time is None:
        errors.append("end_time must be a time like HH:MM.")

    if start_time and end_time and end_time <= start_time:
        errors.append("end_time must be after start_time.")

    if not room:
        errors.append("room is required.")

    if not errors:
        cleaned = {
            "course": course,
            "lecturer": lecturer,
            "day": day,
            "start_time": start_time,
            "end_time": end_time,
            "room": room,
        }

    return errors, cleaned


def timetable_to_dict(entry):
    return {
        "id": entry.id,
        "course": entry.course_id,
        "course_code": entry.course.code,
        "course_name": entry.course.name,
        "lecturer": entry.lecturer_id,
        "lecturer_username": entry.lecturer.username,
        "day": entry.day,
        "start_time": entry.start_time.strftime("%H:%M"),
        "end_time": entry.end_time.strftime("%H:%M"),
        "room": entry.room,
    }


def timetable_queryset():
    return Timetable.objects.select_related("course", "lecturer").order_by("day", "start_time")


def subject_to_dict(subject):
    return {
        "id": subject.id,
        "name": subject.name,
        "code": subject.code,
        "course": subject.course_id,
        "course_name": subject.course.name,
        "course_code": subject.course.code,
    }


def enrollment_to_dict(enrollment):
    student = enrollment.student
    profile = getattr(student, "student", None)
    return {
        "id": enrollment.id,
        "student": student.id,
        "student_username": student.username,
        "student_name": student.get_full_name() or student.username,
        "reg_number": profile.reg_number if profile else None,
        "course": enrollment.course_id,
        "course_name": enrollment.course.name,
        "course_code": enrollment.course.code,
    }


def course_to_dict(course):
    return {
        "id": course.id,
        "name": course.name,
        "code": course.code,
        "department": course.department_id,
        "department_name": course.department.name,
    }


def classroom_to_dict(classroom):
    return {
        "id": classroom.id,
        "room_name": classroom.room_name,
        "room_number": classroom.room_number,
        "latitude": classroom.latitude,
        "longitude": classroom.longitude,
        "radius_meters": classroom.radius_meters,
    }


@api_view(["GET", "POST"])
@permission_classes([IsAdminOrStaff | IsLecturer])
def course_list_create(request):
    if request.method == "GET":
        courses = Course.objects.select_related("department").all()
        return Response([course_to_dict(course) for course in courses])

    if not is_staff_or_superuser(request.user):
        return Response(
            {"detail": "Only admins can create courses."},
            status=status.HTTP_403_FORBIDDEN,
        )

    course = Course.objects.create(
        name=request.data.get("name"),
        code=request.data.get("code"),
        department_id=request.data.get("department_id"),
    )
    return Response({"message": "Course created", "id": course.id}, status=201)


@api_view(["GET", "PATCH", "DELETE"])
@permission_classes([IsAdminOrStaff | IsLecturer])
def course_detail(request, pk):
    course = Course.objects.filter(id=pk).select_related("department").prefetch_related("subjects").first()
    if not course:
        return Response({"detail": "Course not found."}, status=status.HTTP_404_NOT_FOUND)

    if request.method == "GET":
        data = course_to_dict(course)
        data["subjects"] = [
            {"id": subject.id, "name": subject.name, "code": subject.code}
            for subject in course.subjects.all()
        ]
        return Response(data)

    if not is_staff_or_superuser(request.user):
        return Response(
            {"detail": "Only admins can modify courses."},
            status=status.HTTP_403_FORBIDDEN,
        )

    if request.method == "PATCH":
        if "name" in request.data:
            course.name = request.data["name"]
        if "code" in request.data:
            course.code = request.data["code"]
        if "department_id" in request.data:
            course.department_id = request.data["department_id"]
        course.save()
        return Response(course_to_dict(course))

    course.delete()
    return Response({"message": "Course deleted"}, status=status.HTTP_204_NO_CONTENT)


@api_view(["GET", "POST"])
@permission_classes([IsAdminOrStaff | IsLecturer])
def classroom_list_create(request):
    if request.method == "GET":
        return Response([classroom_to_dict(c) for c in Classroom.objects.all()])

    if not is_staff_or_superuser(request.user):
        return Response(
            {"detail": "Only admins can create classrooms."},
            status=status.HTTP_403_FORBIDDEN,
        )

    classroom = Classroom.objects.create(
        room_name=request.data.get("room_name"),
        room_number=request.data.get("room_number"),
        latitude=request.data.get("latitude"),
        longitude=request.data.get("longitude"),
        radius_meters=request.data.get("radius_meters", 20),
    )
    return Response(classroom_to_dict(classroom), status=201)


@api_view(["GET", "PATCH", "DELETE"])
@permission_classes([IsAdminOrStaff | IsLecturer])
def classroom_detail(request, pk):
    classroom = Classroom.objects.filter(id=pk).first()
    if not classroom:
        return Response({"detail": "Classroom not found."}, status=status.HTTP_404_NOT_FOUND)

    if request.method == "GET":
        return Response(classroom_to_dict(classroom))

    if not is_staff_or_superuser(request.user):
        return Response(
            {"detail": "Only admins can modify classrooms."},
            status=status.HTTP_403_FORBIDDEN,
        )

    if request.method == "PATCH":
        for field in ("room_name", "room_number", "latitude", "longitude", "radius_meters"):
            if field in request.data:
                setattr(classroom, field, request.data[field])
        classroom.save()
        return Response(classroom_to_dict(classroom))

    classroom.delete()
    return Response({"message": "Classroom deleted"}, status=status.HTTP_204_NO_CONTENT)


@api_view(["POST"])
@permission_classes([IsAdminOrStaff])
def assign_student_to_course(request):
    student_id = request.data.get("student_id")
    course_id = request.data.get("course_id")

    student = User.objects.filter(id=student_id, groups__name__iexact="Student").first()
    if not student:
        return Response({"detail": "Student user not found."}, status=status.HTTP_404_NOT_FOUND)

    course = Course.objects.filter(id=course_id).first()
    if not course:
        return Response({"detail": "Course not found."}, status=status.HTTP_404_NOT_FOUND)

    StudentCourse.objects.get_or_create(student=student, course=course)
    return Response({"message": "Student assigned successfully"})


@api_view(["POST"])
@permission_classes([IsAdminOrStaff])
def assign_lecturer_to_course(request):
    lecturer_id = request.data.get("lecturer_id")
    course_id = request.data.get("course_id")

    lecturer = User.objects.filter(id=lecturer_id, groups__name__iexact="Lecturer").first()
    if not lecturer:
        return Response({"detail": "Lecturer user not found."}, status=status.HTTP_404_NOT_FOUND)

    course = Course.objects.filter(id=course_id).first()
    if not course:
        return Response({"detail": "Course not found."}, status=status.HTTP_404_NOT_FOUND)

    LecturerCourse.objects.get_or_create(lecturer=lecturer, course=course)
    return Response({"message": "Lecturer assigned successfully"})


@api_view(["GET", "POST"])
@permission_classes([IsAdminOrStaff])
def timetable_list_create(request):
    if request.method == "GET":
        entries = list(timetable_queryset())

        course_id = request.query_params.get("course_id")
        if course_id and course_id.isdigit():
            entries = [entry for entry in entries if entry.course_id == int(course_id)]

        day = request.query_params.get("day")
        if day:
            day_code = parse_day_code(day)
            entries = [entry for entry in entries if entry.day == day_code] if day_code else []

        entries.sort(key=lambda entry: (DAY_ORDER.get(entry.day, 99), entry.start_time))
        return Response([timetable_to_dict(entry) for entry in entries])

    errors, cleaned = validate_timetable_data(request.data)
    if errors:
        return Response({"errors": [{"row": None, "errors": errors}]}, status=status.HTTP_400_BAD_REQUEST)

    entry = Timetable.objects.create(**cleaned)
    return Response(timetable_to_dict(entry), status=status.HTTP_201_CREATED)


@api_view(["GET", "PATCH", "DELETE"])
@permission_classes([IsAdminOrStaff])
def timetable_detail(request, pk):
    entry = Timetable.objects.select_related("course", "lecturer").filter(id=pk).first()
    if not entry:
        return Response({"detail": "Timetable entry not found."}, status=status.HTTP_404_NOT_FOUND)

    if request.method == "GET":
        return Response(timetable_to_dict(entry))

    if request.method == "PATCH":
        merged = {
            "course_code": request.data.get("course_code", entry.course.code),
            "lecturer": request.data.get("lecturer", entry.lecturer.username),
            "day": request.data.get("day", entry.day),
            "start_time": request.data.get("start_time", entry.start_time.strftime("%H:%M")),
            "end_time": request.data.get("end_time", entry.end_time.strftime("%H:%M")),
            "room": request.data.get("room", entry.room),
        }
        errors, cleaned = validate_timetable_data(merged)
        if errors:
            return Response({"errors": [{"row": None, "errors": errors}]}, status=status.HTTP_400_BAD_REQUEST)

        old_start = entry.start_time
        old_end = entry.end_time

        for field, value in cleaned.items():
            setattr(entry, field, value)
        entry.save()

        if entry.start_time != old_start or entry.end_time != old_end:
            from attendance.services import record_timetable_change

            record_timetable_change(
                entry,
                old_start,
                old_end,
                entry.start_time,
                entry.end_time,
                changed_by=request.user,
            )

        return Response(timetable_to_dict(entry))

    entry.delete()
    return Response({"message": "Timetable entry deleted"})


@api_view(["POST"])
@permission_classes([IsAdminOrStaff])
def timetable_upload(request):
    if "file" not in request.FILES:
        return Response(
            {"errors": [{"row": None, "errors": ["Excel file is required."]}]},
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        workbook = load_workbook(BytesIO(request.FILES["file"].read()), data_only=True)
    except Exception:
        return Response(
            {"errors": [{"row": None, "errors": ["Could not read the file. Use a valid .xlsx spreadsheet."]}]},
            status=status.HTTP_400_BAD_REQUEST,
        )

    sheet = workbook.active
    row_iterator = sheet.iter_rows(values_only=True)

    try:
        header_row = next(row_iterator)
    except StopIteration:
        return Response(
            {"errors": [{"row": None, "errors": ["The spreadsheet is empty."]}]},
            status=status.HTTP_400_BAD_REQUEST,
        )

    headers = [cell_text(column).lower() for column in header_row]
    missing_columns = [column for column in TIMETABLE_COLUMNS if column not in headers]
    if missing_columns:
        return Response(
            {"errors": [{"row": None, "errors": [f"Missing columns: {', '.join(missing_columns)}"]}]},
            status=status.HTTP_400_BAD_REQUEST,
        )

    data_rows = []
    for index, values in enumerate(row_iterator, start=2):
        raw = dict(zip(headers, values))
        if not any(cell_text(value) for value in raw.values()):
            continue
        data_rows.append((index, raw))

    if not data_rows:
        return Response(
            {"errors": [{"row": None, "errors": ["No data rows found in the spreadsheet."]}]},
            status=status.HTTP_400_BAD_REQUEST,
        )

    result = {
        "message": "Timetable processed",
        "total_rows": len(data_rows),
        "created_count": 0,
        "updated_count": 0,
        "skipped_count": 0,
        "errors": [],
    }

    seen = set()

    with transaction.atomic():
        for index, raw in data_rows:
            errors, cleaned = validate_timetable_data(raw)

            if not errors:
                key = (cleaned["course"].id, cleaned["lecturer"].id, cleaned["day"], cleaned["start_time"])
                if key in seen:
                    errors.append("Duplicate row in this file.")
                else:
                    seen.add(key)

            if errors:
                result["skipped_count"] += 1
                result["errors"].append({"row": index, "errors": errors})
                continue

            existing = Timetable.objects.filter(
                course=cleaned["course"],
                lecturer=cleaned["lecturer"],
                day=cleaned["day"],
                start_time=cleaned["start_time"],
            ).first()

            if existing:
                for field, value in cleaned.items():
                    setattr(existing, field, value)
                existing.save()
                result["updated_count"] += 1
            else:
                Timetable.objects.create(**cleaned)
                result["created_count"] += 1

    return Response(result)


# ==============================
# SUBJECTS
# ==============================

@api_view(["GET", "POST"])
@permission_classes([IsAdminOrStaff | IsLecturer])
def subject_list_create(request):
    if request.method == "GET":
        subjects = Subject.objects.select_related("course").order_by("course__name", "name")
        course_id = request.query_params.get("course_id")
        if course_id and course_id.isdigit():
            subjects = subjects.filter(course_id=int(course_id))
        return Response([subject_to_dict(subject) for subject in subjects])

    if not is_staff_or_superuser(request.user):
        return Response(
            {"detail": "Only admins can create subjects."},
            status=status.HTTP_403_FORBIDDEN,
        )

    name = cell_text(request.data.get("name"))
    code = cell_text(request.data.get("code"))
    course_id = request.data.get("course_id")

    errors = []
    if not name:
        errors.append("name is required.")
    if not code:
        errors.append("code is required.")

    course = None
    if not course_id:
        errors.append("course_id is required.")
    else:
        course = Course.objects.filter(id=course_id).first()
        if not course:
            errors.append("course_id does not match an existing course.")

    if course and code and Subject.objects.filter(course=course, code__iexact=code).exists():
        errors.append(f"Subject code '{code}' already exists for this course.")

    if errors:
        return Response({"errors": [{"row": None, "errors": errors}]}, status=status.HTTP_400_BAD_REQUEST)

    subject = Subject.objects.create(name=name, code=code, course=course)
    return Response(subject_to_dict(subject), status=status.HTTP_201_CREATED)


@api_view(["GET", "PATCH", "DELETE"])
@permission_classes([IsAdminOrStaff | IsLecturer])
def subject_detail(request, pk):
    subject = Subject.objects.select_related("course").filter(id=pk).first()
    if not subject:
        return Response({"detail": "Subject not found."}, status=status.HTTP_404_NOT_FOUND)

    if request.method == "GET":
        return Response(subject_to_dict(subject))

    if not is_staff_or_superuser(request.user):
        return Response(
            {"detail": "Only admins can modify subjects."},
            status=status.HTTP_403_FORBIDDEN,
        )

    if request.method == "PATCH":
        name = cell_text(request.data.get("name", subject.name))
        code = cell_text(request.data.get("code", subject.code))
        course_id = request.data.get("course_id", subject.course_id)

        errors = []
        if not name:
            errors.append("name is required.")
        if not code:
            errors.append("code is required.")

        course = Course.objects.filter(id=course_id).first()
        if not course:
            errors.append("course_id does not match an existing course.")
        elif Subject.objects.filter(course=course, code__iexact=code).exclude(id=subject.id).exists():
            errors.append(f"Subject code '{code}' already exists for this course.")

        if errors:
            return Response({"errors": [{"row": None, "errors": errors}]}, status=status.HTTP_400_BAD_REQUEST)

        subject.name = name
        subject.code = code
        subject.course = course
        subject.save()
        return Response(subject_to_dict(subject))

    subject.delete()
    return Response({"message": "Subject deleted"}, status=status.HTTP_200_OK)


# ==============================
# STUDENT ENROLLMENT
# ==============================

@api_view(["GET", "POST"])
@permission_classes([IsAdminOrStaff])
def enrollment_list_create(request):
    if request.method == "GET":
        enrollments = StudentCourse.objects.select_related(
            "student", "course", "student__student"
        ).order_by("course__name", "student__username")

        course_id = request.query_params.get("course_id")
        if course_id and course_id.isdigit():
            enrollments = enrollments.filter(course_id=int(course_id))

        student_id = request.query_params.get("student_id")
        if student_id and student_id.isdigit():
            enrollments = enrollments.filter(student_id=int(student_id))

        return Response([enrollment_to_dict(item) for item in enrollments])

    student_id = request.data.get("student_id")
    course_id = request.data.get("course_id")

    student = User.objects.filter(id=student_id, groups__name__iexact="Student").first()
    if not student:
        return Response({"detail": "Student user not found."}, status=status.HTTP_404_NOT_FOUND)

    course = Course.objects.filter(id=course_id).first()
    if not course:
        return Response({"detail": "Course not found."}, status=status.HTTP_404_NOT_FOUND)

    enrollment, created = StudentCourse.objects.get_or_create(student=student, course=course)
    if not created:
        return Response(
            {"detail": "Student is already enrolled in this course."},
            status=status.HTTP_400_BAD_REQUEST,
        )
    return Response(enrollment_to_dict(enrollment), status=status.HTTP_201_CREATED)


@api_view(["DELETE"])
@permission_classes([IsAdminOrStaff])
def enrollment_detail(request, pk):
    enrollment = StudentCourse.objects.select_related(
        "student", "course", "student__student"
    ).filter(id=pk).first()
    if not enrollment:
        return Response({"detail": "Enrollment not found."}, status=status.HTTP_404_NOT_FOUND)

    enrollment.delete()
    return Response({"message": "Enrollment removed"}, status=status.HTTP_200_OK)
